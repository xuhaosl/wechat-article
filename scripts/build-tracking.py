#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build-tracking.py — 扫 output 目录重建文章跟踪表（全自动，零人工列）

2026-10-20 用户定：跟踪表**不记发布状态、不记阅读量**，只客观反映「产出了什么」。
所有列都由文件本身算出来，人工一格都不用填。

列（全部自动）：
  序号 / 日期 / 系列 / 标题 / 汉字数 / 配图 / 封面 / 间隔(天) / 文件

「汉字数」的口径是**去标签后只数汉字**（不含代码块、不含标点、不含英文）。
注意它和 `layout-check.py` 报的「净字数」不是一个东西——后者是全部 `<p>` 文本的
字符总长（含标点英文）。两个数各有用途，别拿来互相校验。

为什么不记「是否已发」「阅读量」：
  · 运营数据是聚合级（tendency/activity/user 都是天或系列），反推不出篇级发布状态
  · 篇级阅读量运营报告里也没有，只有汇总
  · 于是这两个字段只能人工记 —— 而人工列一定会停更（旧表 104 条状态全填「已完成」就是证据）
  真要查篇级数据，去公众号后台导出；表只负责「写了什么、哪天写的」。

用法：
  python build-tracking.py                # 重建并覆盖 output\\公众号文章跟踪表.xlsx
  python build-tracking.py --out x.xlsx   # 输出到别处
  python build-tracking.py --dry          # 只打印，不写文件
"""

import argparse
import os
import re
import sys
from datetime import date, datetime

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = os.environ.get("WECHAT_HOME", "/opt/data/wechat")
OUTPUT = os.path.join(ROOT, "output")
COVERS = os.path.join(OUTPUT, "covers")
XLSX = os.path.join(OUTPUT, "公众号文章跟踪表.xlsx")

# 这些子目录不是正文章节
SKIP_DIRS = {"covers", "tietu", "zhihu", "排版方案", "运营分析", "有趣的Docker", "_raw", "_旧稿"}

NAME_RE = re.compile(r"^(\d{2})_(\d{4}-\d{2}-\d{2})_(.+)\.html$", re.I)
# 漏序号的老稿也认（如 DeepSeek 01 存成了 2026-08-13_xxx.html），序号记 「—」并提示补
LOOSE_RE = re.compile(r"^(\d{4}-\d{2}-\d{2})_(.+)\.html$", re.I)

VISIBLE_RATE = re.compile(r"[\u4e00-\u9fff]")   # 中文字数统计口径：只数汉字


def strip_html(raw):
    """去标签取可见文本。"""
    t = re.sub(r"<script.*?</script>", " ", raw, flags=re.S | re.I)
    t = re.sub(r"<style.*?</style>", " ", t, flags=re.S | re.I)
    t = re.sub(r"<pre.*?</pre>", " ", t, flags=re.S | re.I)
    t = re.sub(r"<[^>]+>", " ", t)
    t = re.sub(r"&nbsp;|&#160;", " ", t)
    t = re.sub(r"&[a-zA-Z]+;|&#\d+;", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def scan():
    """扫 output 目录，返回 (文章列表, 疑似非文章列表)。"""
    arts, others = [], []
    if not os.path.isdir(OUTPUT):
        return arts, others

    covers = [f for f in os.listdir(COVERS)] if os.path.isdir(COVERS) else []

    for dirpath, dirnames, filenames in os.walk(OUTPUT):
        dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS]
        for fn in sorted(filenames):
            if not fn.lower().endswith(".html"):
                continue
            full = os.path.join(dirpath, fn)
            rel_dir = os.path.relpath(dirpath, OUTPUT)
            series = "" if rel_dir == "." else rel_dir
            m = NAME_RE.match(fn)
            if not m:
                loose = LOOSE_RE.match(fn)
                if not loose:
                    others.append({
                        "file": os.path.join(series, fn) if series else fn,
                        "reason": "文件名既无序号也无日期",
                        "size": os.path.getsize(full),
                    })
                    continue
                # 有日期没序号：仍是文章，序号留空待补
                nn, d, title = "—", loose.group(1), loose.group(2)
            else:
                nn, d, title = m.group(1), m.group(2), m.group(3)
            try:
                raw = open(full, encoding="utf-8", errors="replace").read()
            except OSError as e:
                others.append({"file": os.path.join(series, fn), "reason": f"读不了：{e}", "size": 0})
                continue
            text = strip_html(raw)
            cn = len(VISIBLE_RATE.findall(text))
            imgs = len(re.findall(r"<img\s", raw, re.I))

            # 封面：covers 里找「系列名_NN_」或「NN_」开头（无序号的老稿无法定位，记为缺）
            has_cover = ""
            if nn != "—":
                for c in covers:
                    if not c.lower().endswith((".png", ".jpg", ".jpeg")):
                        continue
                    if c.startswith(f"{series}_{nn}_") or (not series and c.startswith(f"{nn}_")):
                        has_cover = c
                        break

            arts.append({
                "nn": nn, "date": d, "series": series or "（根目录）",
                "title": title, "cn": cn, "imgs": imgs,
                "cover": "有" if has_cover else "",
                "cover_file": has_cover,
                "file": os.path.join(series, fn) if series else fn,
            })

    arts.sort(key=lambda x: (x["date"], x["series"], x["nn"]))
    return arts, others


def gap_days(arts):
    """全号口径：距上一篇的天数。"""
    prev = None
    for a in arts:
        try:
            cur = datetime.strptime(a["date"], "%Y-%m-%d").date()
        except ValueError:
            a["gap"] = ""
            continue
        if prev is None:
            a["gap"] = ""
        else:
            a["gap"] = (cur - prev).days
        prev = cur
    return arts


def write_xlsx(arts, others, out_path):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[错误] 缺 openpyxl，先 pip install openpyxl")
        return False

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "产出清单"

    head = ["序号", "日期", "系列", "标题", "汉字数", "配图", "封面", "间隔(天)", "文件"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for a in arts:
        ws.append([a["nn"], a["date"], a["series"], a["title"],
                   a["cn"], a["imgs"], a["cover"] or "—",
                   a["gap"] if a["gap"] != "" else "—", a["file"]])

    widths = {"A": 6, "B": 12, "C": 14, "D": 42, "E": 10, "F": 6, "G": 6, "H": 9, "I": 52}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(horizontal="center")
        row[1].alignment = Alignment(horizontal="center")
        row[4].alignment = Alignment(horizontal="right")
        row[5].alignment = Alignment(horizontal="center")
        row[6].alignment = Alignment(horizontal="center")
        row[7].alignment = Alignment(horizontal="center")

    # 第二页：非文章文件，别把报表当稿子数进去
    ws2 = wb.create_sheet("非文章文件")
    ws2.append(["文件", "原因"])
    for c in (1, 2):
        ws2.cell(row=1, column=c).font = Font(bold=True)
    for o in others:
        ws2.append([o["file"], o["reason"]])
    ws2.column_dimensions["A"].width = 56
    ws2.column_dimensions["B"].width = 40

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=XLSX)
    ap.add_argument("--dry", action="store_true", help="只打印，不写文件")
    args = ap.parse_args()

    arts, others = scan()
    gap_days(arts)

    print("=" * 68)
    print("产出清单 · build-tracking")
    print(f"  扫描目录 {OUTPUT}")
    print("=" * 68)
    print(f"  文章 {len(arts)} 篇 · 非文章 html {len(others)} 个")

    from collections import Counter
    c = Counter(a["series"] for a in arts)
    print("\n  按系列：")
    for k, v in sorted(c.items(), key=lambda x: -x[1]):
        print(f"    {k:16s} {v:3d} 篇")

    if arts:
        ds = [a["date"] for a in arts]
        print(f"\n  日期跨度：{min(ds)} → {max(ds)}")
        total = sum(a["cn"] for a in arts)
        print(f"  总字数：{total:,} 汉字（篇均 {total // len(arts):,}）")
        no_cover = [a for a in arts if not a["cover"]]
        print(f"  缺封面：{len(no_cover)} 篇")
        for a in no_cover[:8]:
            print(f"    {a['date']} {a['series']}/{a['nn']} {a['title'][:30]}")
        no_nn = [a for a in arts if a["nn"] == "—"]
        if no_nn:
            print(f"  文件名漏序号：{len(no_nn)} 篇（建议补成 NN_日期_标题.html）")
            for a in no_nn:
                print(f"    {a['series']} / {a['file']}")
        gaps = [a["gap"] for a in arts if isinstance(a["gap"], int)]
        if gaps:
            print(f"  更新间隔：中位 {sorted(gaps)[len(gaps)//2]} 天 · 最长 {max(gaps)} 天")

    if others:
        print(f"\n  非文章 html（不进清单）：")
        for o in others:
            print(f"    {o['file']}")

    if args.dry:
        print("\n[--dry] 未写文件")
        return 0

    if write_xlsx(arts, others, args.out):
        print(f"\n✅ 已写入 {args.out}")
        return 0
    return 2


if __name__ == "__main__":
    sys.exit(main())
